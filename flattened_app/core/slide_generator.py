from pptx import Presentation
from pptx.util import Inches, Pt

def generate_slides(summaries: list[str], out_path: str):
    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Lecture Slides"
    subtitle.text = "Auto-generated from PDF"

    bullet_layout = prs.slide_layouts[1]
    for idx, summary in enumerate(summaries, start=1):
        slide = prs.slides.add_slide(bullet_layout)
        shapes = slide.shapes
        shapes.title.text = f"Section {idx}"
        body_shape = shapes.placeholders[1]
        tf = body_shape.text_frame
        for sentence in summary.split(". "):
            p = tf.add_paragraph()
            p.text = sentence.strip()
            p.font.size = Pt(18)

    prs.save(out_path)
